# -*- coding: utf-8 -*-
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

path = Path(r"c:\Users\Nicolas\Downloads\0.0 - Modelo de Cierre de Negocio-UARM  mejoras.xlsx")
if not path.exists():
    print("No existe:", path)
    sys.exit(1)

wb = openpyxl.load_workbook(path, data_only=True)
wb_f = openpyxl.load_workbook(path, data_only=False)

print("=== ARCHIVO ===")
print("Hojas:", wb.sheetnames)
print()

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    ws_f = wb_f[sheet_name]
    print(f"=== HOJA: {sheet_name} ===")
    print(f"Dimensiones: {ws.max_row} x {ws.max_column}")
    merged = list(ws_f.merged_cells.ranges)
    print(f"Merged: {len(merged)}")
    shown = 0
    for row in range(1, min(ws.max_row + 1, 80)):
        cells = []
        for col in range(1, min(ws.max_column + 1, 25)):
            val = ws.cell(row=row, column=col).value
            if val is not None and str(val).strip() != "":
                cells.append(f"{get_column_letter(col)}={repr(val)[:60]}")
        if cells:
            print(f"  F{row}: " + " | ".join(cells))
            shown += 1
            if shown >= 55:
                print("  ... truncado")
                break
    print()
