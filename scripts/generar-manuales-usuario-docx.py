# -*- coding: utf-8 -*-
"""
Genera un .docx por módulo desde docs/manuales-usuario/*.md
(o docs/manuales-portal/*.md si no hay carpeta usuario).

Salida: ~/Downloads/Manuales-Portal-Prayaga/
Requiere: pip install python-docx
"""
from __future__ import annotations

import importlib.util
import re
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SOURCE_DIR = ROOT / "docs" / "manuales-usuario"
FALLBACK_DIR = ROOT / "docs" / "manuales-portal"
OUT_DIR = Path.home() / "Downloads" / "Manuales-Portal-Prayaga"
INFORME_SCRIPT = ROOT / "scripts" / "generar-informe-portal-docx.py"


def _load_renderer():
    spec = importlib.util.spec_from_file_location("informe_docx", INFORME_SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod.render_markdown


render_markdown = _load_renderer()


def add_cover(doc: Document, titulo: str, codigo: str, audiencia: str):
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("PRAYAGA")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E)

    doc.add_paragraph()
    h = doc.add_heading(f"Manual de usuario — {titulo}", level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Portal Prayaga Interno\n").bold = True
    meta.add_run(f"Código: {codigo}  |  Versión: 1.0  |  {date.today().strftime('%d/%m/%Y')}\n")
    meta.add_run(f"Audiencia: {audiencia}")

    doc.add_paragraph()
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.add_run(
        "Documento de uso operativo. En el portal, use «Guía de uso» en la parte superior del módulo "
        "para consulta rápida."
    ).italic = True
    doc.add_page_break()


def md_files():
    if SOURCE_DIR.is_dir():
        files = sorted(SOURCE_DIR.glob("*.md"))
        files = [f for f in files if f.name.lower() != "readme.md"]
        if files:
            return files
    if FALLBACK_DIR.is_dir():
        return sorted(
            f for f in FALLBACK_DIR.glob("*.md") if f.name.lower() != "readme.md"
        )
    return []


def titulo_desde_nombre(path: Path) -> str:
    text = path.read_text(encoding="utf-8")
    m = re.match(r"^#\s+(.+?)(?:\s+—|\s+-|$)", text, re.M)
    if m:
        return re.sub(r"\s*—\s*Manual de uso\s*$", "", m.group(1).strip())
    return path.stem.replace("-", " ").title()


def codigo_doc(path: Path) -> str:
    slug = re.sub(r"[^a-z0-9]", "", path.stem.lower())[:10].upper()
    return f"MAN-{slug}-001"


def audiencia_default(text: str) -> str:
    if "Restringido" in text or "RESTRINGIDO" in text:
        return "Usuarios autorizados (administrador o módulo activado en Administración de usuarios)"
    return "Colaboradores con acceso al módulo en el portal"


def build_one(path: Path, out_dir: Path) -> Path:
    md = path.read_text(encoding="utf-8")
    titulo = titulo_desde_nombre(path)
    doc = Document()
    add_cover(doc, titulo, codigo_doc(path), audiencia_default(md))
    render_markdown(doc, md)
    out = out_dir / f"Manual-{path.stem}.docx"
    doc.save(out)
    return out


def main():
    files = md_files()
    if not files:
        print("No hay archivos .md en docs/manuales-usuario ni manuales-portal")
        sys.exit(1)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for f in files:
        out = build_one(f, OUT_DIR)
        print(out)
    print(f"\n{len(files)} manuales en: {OUT_DIR}")


if __name__ == "__main__":
    main()
