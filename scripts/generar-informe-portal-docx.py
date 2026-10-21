# -*- coding: utf-8 -*-
"""Genera Informe-Arquitectura-Portal-Prayaga.docx en Descargas del usuario."""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTPUT = Path.home() / "Downloads" / "Informe-Arquitectura-Portal-Prayaga.docx"
MD_SOURCE = Path(__file__).resolve().parent.parent / "docs" / "Informe-Arquitectura-Portal-Prayaga.md"


def set_cell_shading(cell, fill: str):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        set_cell_shading(hdr[i], "1F4E79")
        for p in hdr[i].paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    doc.add_paragraph()


def add_code_block(doc, text: str):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def parse_inline(text: str):
    """Return list of (text, bold) segments for simple **bold** markup."""
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    segments = []
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            segments.append((part[2:-2], True))
        elif part:
            segments.append((part, False))
    return segments


def add_rich_paragraph(doc, text: str, style=None):
    p = doc.add_paragraph(style=style)
    for seg, bold in parse_inline(text):
        run = p.add_run(seg)
        run.bold = bold
    return p


def render_markdown(doc, md_text: str):
    lines = md_text.splitlines()
    i = 0
    in_code = False
    code_buf = []
    table_buf = []

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            if in_code:
                add_code_block(doc, "\n".join(code_buf))
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if line.startswith("|") and "|" in line[1:]:
            table_buf.append(line)
            i += 1
            continue
        elif table_buf:
            rows_raw = [r for r in table_buf if not re.match(r"^\|[\s\-:|]+\|$", r.strip())]
            if rows_raw:
                parsed = [[c.strip() for c in r.strip("|").split("|")] for r in rows_raw]
                if parsed:
                    add_table(doc, parsed[0], parsed[1:])
            table_buf = []

        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=4)
        elif line.strip() == "---":
            doc.add_paragraph()
        elif line.startswith("- "):
            add_rich_paragraph(doc, line[2:].strip(), style="List Bullet")
        elif line.strip().startswith(">"):
            add_rich_paragraph(doc, line.strip().lstrip(">").strip())
        elif line.strip():
            add_rich_paragraph(doc, line.strip())
        i += 1

    if table_buf:
        rows_raw = [r for r in table_buf if not re.match(r"^\|[\s\-:|]+\|$", r.strip())]
        if rows_raw:
            parsed = [[c.strip() for c in r.strip("|").split("|")] for r in rows_raw]
            if parsed:
                add_table(doc, parsed[0], parsed[1:])


def build_document():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("Informe técnico de arquitectura", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("Portal de Prayaga (gestor-vacaciones)")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = doc.add_paragraph(
        "Repositorio: Nicolas0903/Gestion-de-Vacaciones | "
        "Alcance: análisis basado en código fuente | Octubre 2026"
    )
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    if MD_SOURCE.exists():
        render_markdown(doc, MD_SOURCE.read_text(encoding="utf-8"))
    else:
        doc.add_paragraph(f"No se encontró {MD_SOURCE}. Ejecute desde el repositorio.")

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build_document()
    print(f"Generado: {path}")
