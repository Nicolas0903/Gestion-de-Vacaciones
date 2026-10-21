# -*- coding: utf-8 -*-
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

path = Path(r"c:\Users\Nicolas\Downloads\Costos y Cronograma Jockey Plaza.xlsx")
if not path.exists():
    print("No existe:", path)
    sys.exit(1)

wb = openpyxl.load_workbook(path, data_only=False)
wb_values = openpyxl.load_workbook(path, data_only=True)

print("=== ARCHIVO ===")
print("Nombre:", path.name)
print("Hojas:", wb.sheetnames)
print()

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    ws_v = wb_values[sheet_name]
    print(f"=== HOJA: {sheet_name} ===")
    print(f"Dimensiones: {ws.max_row} filas x {ws.max_column} columnas")
    print(f"Rango usado: {ws.calculate_dimension()}")

    merged = list(ws.merged_cells.ranges)
    print(f"Celdas combinadas: {len(merged)}")
    for m in merged[:20]:
        print(f"  {m}")
    if len(merged) > 20:
        print(f"  ... y {len(merged) - 20} mas")

    print("Contenido (filas con datos, cols A-Q):")
    shown = 0
    for row in range(1, min(ws.max_row + 1, 150)):
        cells = []
        for col in range(1, min(ws.max_column + 1, 18)):
            cell = ws.cell(row=row, column=col)
            val = cell.value
            if val is not None and str(val).strip() != "":
                tag = " [F]" if isinstance(val, str) and val.startswith("=") else ""
                val_v = ws_v.cell(row=row, column=col).value
                cells.append(
                    f"{get_column_letter(col)}{row}={repr(val)[:50]}{tag} calc={repr(val_v)[:30]}"
                )
        if cells:
            print(f"  F{row}:")
            for c in cells:
                print(f"    {c}")
            shown += 1
            if shown >= 50:
                print("  ... (truncado)")
                break
    print()

print("=== FORMULAS (muestra) ===")
formula_count = 0
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                print(f"  {sheet_name}!{cell.coordinate}: {cell.value[:80]}")
                formula_count += 1
                if formula_count >= 30:
                    break
        if formula_count >= 30:
            break
    if formula_count >= 30:
        break
print(f"Total formulas mostradas (max 30): {formula_count}")

print("\n=== RANGOS CON NOMBRE ===")
if wb.defined_names:
    for name in wb.defined_names:
        print(f"  {name}: {wb.defined_names[name].attr_text}")
else:
    print("  (ninguno)")
